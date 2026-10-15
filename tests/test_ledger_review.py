import copy, io, unittest
from openpyxl import load_workbook
from pypdf import PdfReader
import ledger, workflow, server, erp_documents


class ReviewTests(unittest.TestCase):
 def setUp(self):
  self.s=server.empty_state();self.s['accounts']=[{'id':'buyer','name':'테스트 고객'},{'id':'vendor','name':'테스트 공급처'}]
  workflow.normalize(self.s);self.command('erp.initialize',confirmed=True)
 def command(self, action, **data):
  result, content=workflow.apply(self.s,dict(action=action,data=data),'demo',actor='test-owner')
  server.validate_state(self.s);return result,content
 def product(self):
  return self.command('erp.product.save',code='P',name='상자',unit='EA',purchasePrice=100000,salePrice=150000)[0]['recordId']
 def voucher(self,kind='sale',date='2026-09-01',qty=1,price=100000,pid='',**more):
  d=dict(kind=kind,title='테스트',accountId='buyer' if kind=='sale' else 'vendor',date=date,items=[dict(name='상자',productId=pid,quantity=qty,unit='EA',unitPrice=price,vatRate=10)])
  d.update(more);return self.command('erp.voucher.save',**d)[0]['recordId']
 def v(self,key):return ledger.get(self.s,'vouchers',key)
 def fulfill(self,key,qty=1,date='2026-09-30'):
  return self.command('erp.fulfill',id=key,date=date,items=[dict(lineId=self.v(key)['items'][0]['id'],quantity=qty)])[0]['fulfillmentId']
 def bal(self,a,dt='9999-12-31'):
  return sum(r['debit']-r['credit'] for j in self.s['journalEntries'] if j['date']<=dt for r in j['rows'] if r['account']==a)
 def order(self):
  q=self.command('quote.save',accountId='buyer',title='B건',date='2026-09-01',items=[dict(name='상자',unit='EA',quantity=1,unitPrice=100000)])[0]['recordId']
  return self.command('quote.convert',quoteId=q,date='2026-09-01',orderType='BOX')[0]['recordId']

 def test_products_can_be_deleted_atomically_in_bulk(self):
  first=self.command('erp.product.save',code='BULK-1',name='일괄 삭제 1',unit='EA',purchasePrice=100,salePrice=200)[0]['recordId']
  second=self.command('erp.product.save',code='BULK-2',name='일괄 삭제 2',unit='EA',purchasePrice=300,salePrice=400)[0]['recordId']
  self.command('erp.stock.opening',productId=first,date='2026-09-01',quantity=2,value=200)
  voucher=self.command('erp.voucher.save',kind='sale',accountId='buyer',title='스냅샷 보존',date='2026-09-02',items=[dict(name='일괄 삭제 1',productId=first,quantity=1,unit='EA',unitPrice=200,vatRate=10),dict(name='일괄 삭제 2',productId=second,quantity=1,unit='EA',unitPrice=400,vatRate=10)])[0]['recordId']
  result,_=self.command('erp.product.bulk-delete',ids=[first,second])
  self.assertEqual(result,dict(recordIds=[first,second],count=2))
  self.assertFalse(any(p['id'] in (first,second) for p in self.s['products']))
  self.assertFalse(any(m['productId'] in (first,second) for m in self.s['stockMoves']))
  self.assertEqual([row['productId'] for row in self.v(voucher)['items']],['',''])
  self.assertEqual([row['name'] for row in self.v(voucher)['items']],['일괄 삭제 1','일괄 삭제 2'])
  self.assertEqual(sum(a['action']=='품목·재고 이력 삭제' and a['entityId'] in (first,second) for a in self.s['erpAudit']),2)

 def test_bulk_product_delete_rolls_back_when_any_id_is_invalid(self):
  product=self.product();before=copy.deepcopy(self.s)
  with self.assertRaisesRegex(ValueError,'연결된 기록'):
   self.command('erp.product.bulk-delete',ids=[product,'missing-product'])
  self.assertEqual(self.s,before)

 def test_late_supply_before_confirmation_and_voucher_date_posts_to_original_month(self):
  p=self.product();self.command('erp.stock.opening',productId=p,date='2026-09-01',quantity=3,value=300000)
  v=self.voucher(date='2026-10-02',pid=p,price=150000)
  self.command('erp.voucher.confirm',id=v,date='2026-10-02');self.fulfill(v)
  self.assertEqual(self.bal('revenue','2026-09-30'),-150000)
  self.assertEqual(self.bal('cogs','2026-09-30'),100000)
  self.assertEqual(self.bal('unbilledSales','2026-09-30'),150000)
  self.assertEqual(self.bal('ar','2026-09-30'),0)
  self.assertEqual(self.bal('unbilledSales'),0);self.assertEqual(self.bal('deferredSales'),0)
  self.assertEqual(ledger.stock(self.s,p),dict(quantity=2,value=200000))

 def test_backdated_receipt_revalues_later_cogs(self):
  p=self.product();self.command('erp.stock.opening',productId=p,date='2026-09-01',quantity=1,value=100000)
  sale=self.voucher(pid=p,price=250000);self.command('erp.voucher.confirm',id=sale,date='2026-09-01');self.fulfill(sale,date='2026-09-20')
  buy=self.voucher('purchase',date='2026-10-01',pid=p,price=200000);self.command('erp.voucher.confirm',id=buy,date='2026-10-01');self.fulfill(buy,date='2026-09-10')
  self.assertEqual(self.bal('cogs'),150000)
  self.assertEqual(ledger.stock(self.s,p),dict(quantity=1,value=150000))
  self.assertEqual(self.bal('accruedPurchase','2026-09-30'),-200000)
  self.assertEqual(self.bal('accruedPurchase'),0)

 def test_late_supply_revalues_next_sale_and_keeps_balanced_periods(self):
  p=self.product();self.command('erp.stock.opening',productId=p,date='2026-09-01',quantity=2,value=200000)
  buy=self.voucher('purchase',pid=p,qty=2,price=200000);self.fulfill(buy,2,'2026-09-20')
  sale=self.voucher(pid=p);self.command('erp.voucher.confirm',id=sale,date='2026-10-01');self.fulfill(sale,date='2026-09-25')
  late=self.voucher(pid=p);self.command('erp.voucher.confirm',id=late,date='2026-10-02');f=self.fulfill(late,date='2026-09-10')
  self.assertEqual(ledger.get(self.s,'fulfillments',f)['cost'],100000)
  self.assertEqual(self.bal('cogs'),266667)
  self.assertEqual(ledger.stock(self.s,p)['value'],333333)
  self.assertEqual(self.bal('inventory'),333333)

 def test_historical_shortage_is_advisory(self):
  p=self.product();self.command('erp.stock.opening',productId=p,date='2026-10-01',quantity=2,value=200000)
  v=self.voucher(pid=p);f=self.fulfill(v,date='2026-09-30')
  self.assertEqual(ledger.stock(self.s,p,'2026-09-30'),dict(quantity=-1,value=-100000))
  self.assertEqual(ledger.stock(self.s,p),dict(quantity=1,value=100000))
  self.assertTrue(ledger.get(self.s,'fulfillments',f)['costProvisional'])

 def test_duplicate_order_voucher_all_save_routes_and_approval_number(self):
  o=self.order();v=self.command('erp.order.voucher',orderId=o,date='2026-09-01')[0]['recordId']
  self.assertEqual(self.command('erp.order.voucher',orderId=o,date='2026-09-01')[0]['recordId'],v)
  with self.assertRaisesRegex(ValueError,'이미 매출 전표'):self.voucher(orderId=o)
  self.command('erp.voucher.confirm',id=v,date='2026-09-01')
  self.command('erp.voucher.approval',id=v,approvalNumber='TEST-20260901')
  v2=self.voucher();self.command('erp.voucher.confirm',id=v2,date='2026-09-01')
  with self.assertRaisesRegex(ValueError,'이미 등록된 승인번호'):self.command('erp.voucher.approval',id=v2,approvalNumber='TEST20260901')

 def test_optional_evidence_does_not_change_financials_or_exports_and_keeps_actor(self):
  v=self.voucher();self.command('erp.voucher.evidence',id=v,taxEvidence={})
  self.command('erp.voucher.confirm',id=v,date='2026-09-01');before=copy.deepcopy(self.s['journalEntries'])
  m=dict(evidenceType='전자세금계산서',taxCategory='과세',writtenDate='2026-09-01',issuedDate='2026-10-02',sourceNet=100000,sourceVat=10000,sourceGross=110000,deductibleVat=0,nondeductionReason='검토 메모',accountantMemo='INTERNAL-TAX-ONLY')
  self.command('erp.voucher.evidence',id=v,taxEvidence=m)
  self.assertEqual(self.s['journalEntries'],before);self.assertEqual(self.v(v)['taxEvidence'],m)
  self.assertEqual(self.s['erpAudit'][-1]['actorUserId'],'test-owner')
  snap=erp_documents.snapshot(self.s,'voucher',v);self.assertNotIn('taxEvidence',snap)
  with self.assertRaises(ValueError):self.command('erp.voucher.evidence',id=v,taxEvidence={'writtenDate':'bad'})

 def test_order_difference_preserves_original_and_xlsx_parity(self):
  o=self.order();original=copy.deepcopy(ledger.get(self.s,'orders',o))
  v=self.command('erp.order.voucher',orderId=o,date='2026-09-01')[0]['recordId']
  self.voucher(id=v,orderId=o,items=[dict(name='상자',quantity=1,unit='EA',unitPrice=100000,vatRate=0)],orderDifferenceReason='면세 여부 확인 중')
  self.assertEqual(self.v(v)['amount'],100000);self.assertEqual(ledger.get(self.s,'orders',o),original)
  r=erp_documents.snapshot(self.s,'voucher',v)
  wb=load_workbook(io.BytesIO(erp_documents.render_xlsx('voucher',r)));ws=wb.active
  self.assertEqual(ws['E13'].value,'단위');self.assertEqual(ws['E14'].value,'EA')
  self.assertIsNone(ws['A3'].value)
  from pdf_documents import render_pdf
  pdf=PdfReader(io.BytesIO(render_pdf('delivery',r)));txt='\n'.join(p.extract_text() for p in pdf.pages)
  self.assertNotIn(original['number'],txt);self.assertNotIn('연결 주문',txt);self.assertIn('EA',txt)
  self.assertNotIn('면세 여부 확인 중',txt)

 def test_unknown_cost_is_distinct_from_explicit_zero(self):
  unknown=self.voucher();self.assertFalse(self.v(unknown)['items'][0]['costKnown'])
  zero=self.voucher(items=[dict(name='무상 자재',quantity=1,unitPrice=100000,costUnit=0,costKnown=True)])
  self.assertTrue(self.v(zero)['items'][0]['costKnown'])

 def test_replay_is_idempotent_and_preserves_payment_and_legacy_postings(self):
  from ledger_replay import rebuild
  self.s['journalEntries'].append(dict(id='legacy-test',sourceId='legacy-source',date='2026-08-01',memo='이관',rows=[dict(account='cash',debit=100,credit=0),dict(account='equity',debit=0,credit=100)]))
  v=self.voucher();self.command('erp.voucher.confirm',id=v,date='2026-10-02')
  self.command('erp.payment.save',direction='in',accountId='buyer',date='2026-10-03',amount=50000,voucherId=v)
  self.fulfill(v,date='2026-09-30')
  before=copy.deepcopy(self.s)
  rebuild(self.s);self.assertEqual(self.s,before)
  self.assertEqual(self.bal('cash'),50100);self.assertEqual(self.bal('ar'),60000)
  self.assertTrue(any(j['id']=='legacy-test' for j in self.s['journalEntries']))

 def test_backdated_receipt_revalues_material_use_expense(self):
  p=self.product();self.command('erp.stock.opening',productId=p,date='2026-09-01',quantity=2,value=200000)
  self.command('erp.stock.use',productId=p,date='2026-09-20',quantity=1,reason='자재 투입')
  b=self.voucher('purchase',pid=p,price=200000);self.fulfill(b,date='2026-09-10')
  self.assertEqual(self.bal('expense'),133333)
  self.assertEqual(ledger.stock(self.s,p),dict(quantity=2,value=266667))

 def test_delivery_allowed_even_when_later_cancelled_receipt_leaves_negative_stock(self):
  p=self.product();buy=self.voucher('purchase',pid=p,price=100000)
  f=self.fulfill(buy,date='2026-09-10')
  self.command('erp.fulfillment.void',id=f,date='2026-09-30',reason='오입력')
  sell=self.voucher(pid=p);self.fulfill(sell,date='2026-09-20')
  self.assertEqual(ledger.stock(self.s,p),dict(quantity=-1,value=-100000))
  self.assertEqual(self.bal('inventory'),-100000)

 def test_missing_receipt_can_be_backfilled_to_replace_provisional_cost(self):
  p=self.product();sell=self.voucher(pid=p,qty=2);f=self.fulfill(sell,2,'2026-09-20')
  self.assertEqual(ledger.get(self.s,'fulfillments',f)['cost'],200000)
  buy=self.voucher('purchase',pid=p,qty=3,price=120000);self.fulfill(buy,3,'2026-09-10')
  f=ledger.get(self.s,'fulfillments',f)
  self.assertEqual(f['cost'],240000);self.assertFalse(f['costProvisional'])
  self.assertEqual(ledger.stock(self.s,p),dict(quantity=1,value=120000))
  self.assertEqual(self.bal('inventory'),120000);self.assertEqual(self.bal('cogs'),240000)

 def test_partial_stock_and_repeated_negative_deliveries_keep_nonnegative_cost(self):
  p=self.product();self.command('erp.stock.opening',productId=p,date='2026-09-01',quantity=1,value=100001)
  v=self.voucher(pid=p,qty=3);self.fulfill(v,2,'2026-09-10');self.fulfill(v,1,'2026-09-11')
  self.assertEqual(ledger.stock(self.s,p),dict(quantity=-2,value=-200002))
  self.assertEqual(self.bal('cogs'),300003)
  self.assertTrue(all(f['costProvisional'] for f in self.s['fulfillments']))
  buy=self.voucher('purchase',pid=p,qty=3,price=50000);self.fulfill(buy,3,'2026-09-12')
  # Positive quantity with a negative provisional valuation must never create
  # negative costs or negative catalog cost inputs on the next sale.
  next_sale=self.voucher(pid=p);self.assertGreaterEqual(self.v(next_sale)['items'][0]['costUnit'],0)
  self.fulfill(next_sale,date='2026-09-13')
  self.assertEqual(self.s['fulfillments'][-1]['cost'],50000)
  self.assertTrue(self.s['fulfillments'][-1]['costProvisional'])

 def test_unknown_cost_zero_stock_and_cancellation_are_replay_stable(self):
  from ledger_replay import rebuild
  p=self.command('erp.product.save',code='NO-COST',name='미입력 품목')[0]['recordId']
  v=self.voucher(pid=p,qty=.5);f=self.fulfill(v,.5)
  self.assertEqual(self.bal('cogs'),0);self.assertTrue(self.s['fulfillments'][-1]['costProvisional'])
  before=copy.deepcopy(self.s);rebuild(self.s);self.assertEqual(self.s,before)
  self.command('erp.fulfillment.void',id=f,date='2026-10-01',reason='정정')
  self.assertEqual(ledger.stock(self.s,p),dict(quantity=0,value=0))
  self.assertFalse(ledger.provisional_stock(self.s,p))
  self.assertEqual(self.bal('revenue'),0)

 def test_catalog_price_edit_does_not_reprice_existing_shortage(self):
  from ledger_replay import rebuild
  p=self.product();v=self.voucher(pid=p);f=self.fulfill(v)
  self.command('erp.product.save',id=p,code='P',name='상자',purchasePrice=999999)
  rebuild(self.s)
  self.assertEqual(ledger.get(self.s,'fulfillments',f)['cost'],100000)

 def test_backdated_delivery_does_not_block_existing_later_material_use(self):
  p=self.product();self.command('erp.stock.opening',productId=p,date='2026-09-01',quantity=1,value=100000)
  self.command('erp.stock.use',productId=p,date='2026-09-20',quantity=1,reason='자재 사용')
  v=self.voucher(pid=p);self.fulfill(v,date='2026-09-10')
  self.assertEqual(ledger.stock(self.s,p)['quantity'],-1)
  self.assertEqual(self.bal('expense'),100000)

 def test_void_after_late_fulfillment_leaves_no_bridge_balance(self):
  v=self.voucher();self.command('erp.voucher.confirm',id=v,date='2026-10-02');f=self.fulfill(v,date='2026-09-30')
  self.command('erp.fulfillment.void',id=f,date='2026-10-03',reason='오입력 정정')
  with self.assertRaisesRegex(ValueError,'취소일 이후'):self.command('erp.voucher.void',id=v,date='2026-10-02',reason='정정')
  self.command('erp.voucher.void',id=v,date='2026-10-03',reason='정정')
  for a in ledger.ACCOUNTS:self.assertEqual(self.bal(a),0,a)
