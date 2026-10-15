import copy,io,unittest
from openpyxl import load_workbook
import ledger,workflow,server

class LedgerTests(unittest.TestCase):
 def setUp(self):
  self.s=server.empty_state();self.s['accounts']=[dict(id='supplier',name='볼트 서플라이'),dict(id='customer',name='세이프 스튜디오')]
  workflow.normalize(self.s);self.run('erp.initialize',confirmed=True)
 def run(self,action=None,**data):
  if action is None or not isinstance(action,str):return super().run(action)
  result,content=workflow.apply(self.s,dict(action=action,data=data),'demo');server.validate_state(self.s);return result,content
 def product(self):
  r,_=self.run('erp.product.save',code='TC80150',name='철제함',spec='800*1500*350*1.6T STEEL',unit='EA',salePrice=150000,purchasePrice=100000,trackStock=True,supplierId='supplier')
  return r['recordId']
 def voucher(self,kind,qty=1,price=150000,pid='',**more):
  d=dict(kind=kind,date='2026-10-01',dueDate='2026-10-05',accountId='customer' if kind=='sale' else 'supplier',title='철제함',items=[dict(name='철제함',productId=pid,quantity=qty,unitPrice=price)],**more)
  r,_=self.run('erp.voucher.save',**d);return ledger.get(self.s,'vouchers',r['recordId'])
 def fulfill(self,v,qty,date='2026-10-01'):
  return self.run('erp.fulfill',id=v['id'],date=date,items=[dict(lineId=v['items'][0]['id'],quantity=qty)])
 def test_purchase_receipt_sale_confirmation_partial_cash_and_fulfillment_are_independent(self):
  pid=self.product();p=self.voucher('purchase',10,100000,pid)
  self.fulfill(p,10);self.run('erp.voucher.confirm',id=p['id'])
  self.assertEqual(ledger.stock(self.s,pid),dict(quantity=10,value=1000000))
  v=self.voucher('sale',2,150000,pid);self.run('erp.voucher.confirm',id=v['id'])
  self.assertEqual(ledger.stock(self.s,pid)['quantity'],10)
  self.assertEqual(ledger.balance(self.s,v),330000)
  self.assertFalse(any(r['account']=='revenue' for j in self.s['journalEntries'] for r in j['rows']))
  self.run('erp.payment.save',direction='in',accountId='customer',date='2026-10-01',amount=100000,voucherId=v['id'])
  self.assertEqual(ledger.balance(self.s,v),230000)
  self.fulfill(v,1);self.fulfill(v,1)
  self.assertEqual(ledger.stock(self.s,pid),dict(quantity=8,value=800000))
  totals={a:sum(r['debit']-r['credit'] for j in self.s['journalEntries'] for r in j['rows'] if r['account']==a) for a in ledger.ACCOUNTS}
  self.assertEqual(totals['revenue'],-300000);self.assertEqual(totals['cogs'],200000)
  self.assertEqual(totals['deferredSales'],0)
  self.run('erp.payment.save',direction='out',accountId='supplier',date='2026-10-01',amount=1100000,voucherId=p['id'])
  self.assertEqual(ledger.balance(self.s,p),0)
 def test_supply_before_invoice_and_prepayment_allocation(self):
  p=self.product();self.run('erp.stock.opening',productId=p,date='2026-09-01',quantity=10,value=1000000)
  pay,_=self.run('erp.payment.save',direction='in',accountId='customer',date='2026-09-20',amount=165000)
  v=self.voucher('sale',1,150000,p);self.fulfill(v,1)
  self.run('erp.voucher.confirm',id=v['id'],date='2026-10-02')
  self.run('erp.payment.allocate',id=pay['recordId'],voucherId=v['id'],amount=165000,date='2026-10-02')
  self.assertEqual(ledger.balance(self.s,v),0)
  with self.assertRaises(ValueError):self.run('erp.payment.allocate',id=pay['recordId'],voucherId=v['id'],amount=1,date='2026-10-02')
 def test_confirm_without_approval_is_idempotent_and_tax_memo_has_no_money_effect(self):
  v=self.voucher('sale');self.run('erp.voucher.confirm',id=v['id']);original=copy.deepcopy(self.s['journalEntries'])
  self.run('erp.voucher.confirm',id=v['id']);self.run('erp.voucher.approval',id=v['id'],approvalNumber='')
  self.run('erp.voucher.approval',id=v['id'],approvalNumber='20261001-12345678-12345678')
  self.assertEqual(self.s['journalEntries'],original)
 def test_insufficient_stock_allowed_but_overfulfillment_rejected(self):
  pid=self.product();v=self.voucher('sale',1,150000,pid)
  self.fulfill(v,1)
  self.assertEqual(ledger.stock(self.s,pid),dict(quantity=-1,value=-100000))
  self.assertTrue(self.s['fulfillments'][-1]['costProvisional'])
  before=copy.deepcopy(self.s)
  with self.assertRaisesRegex(ValueError,'수량을 초과'):self.fulfill(v,1)
  self.assertEqual(self.s,before)
 def test_quote_price_book_snapshot_and_box_number(self):
  pid=self.product();r,_=self.run('quote.save',accountId='customer',date='2026-10-01',title='A건',items=[dict(name='철제함',productId=pid,quantity=2,unitPrice=150000)])
  q=workflow.get(self.s,'quotes',r['recordId']);self.assertEqual(q['items'][0]['estimatedCost'],200000)
  o,_=self.run('quote.convert',quoteId=q['id'],date='2026-10-01',orderType='BOX');order=workflow.get(self.s,'orders',o['recordId']);self.assertEqual(order['number'],'BOX-20261001-001')
  v,_=self.run('erp.order.voucher',orderId=order['id'],date='2026-10-01');again,_=self.run('erp.order.voucher',orderId=order['id'],date='2026-10-01');self.assertEqual(v,again)
  self.run('erp.voucher.confirm',id=v['recordId'])
  self.assertEqual(order['status'],'주문 접수')
 def test_product_delete_preserves_line_snapshots_and_removes_stock_history(self):
  pid=self.product();v=self.voucher('sale',1,150000,pid)
  self.run('erp.product.delete',id=pid)
  self.assertFalse(any(p['id']==pid for p in self.s['products']))
  self.assertEqual(v['items'][0]['productId'],'')
  self.assertEqual(v['items'][0]['name'],'철제함')
  pid=self.product();self.run('erp.stock.opening',productId=pid,date='2026-09-01',quantity=1,value=100000)
  self.run('erp.product.delete',id=pid)
  self.assertFalse(any(p['id']==pid for p in self.s['products']))
  self.assertFalse(any(m['productId']==pid for m in self.s['stockMoves']))
 def test_account_delete_cleans_lightweight_links_and_cascades_business_records(self):
  self.s['accounts'].append(dict(id='unused',name='삭제 거래처'))
  self.s['tasks'].append(dict(id='task-unused',accountId='unused',title='정리',date='2026-10-01',done=False))
  self.s['notes'].append(dict(id='note-unused',accountId='unused',date='2026-10-01',body='정리'))
  pid=self.product();ledger.get(self.s,'products',pid)['supplierId']='unused'
  self.run('erp.account.delete',id='unused')
  self.assertFalse(any(a['id']=='unused' for a in self.s['accounts']))
  self.assertFalse(any(r.get('accountId')=='unused' for r in self.s['tasks']+self.s['notes']))
  self.assertEqual(ledger.get(self.s,'products',pid)['supplierId'],'')
  voucher=self.voucher('sale')
  self.run('erp.account.delete',id='customer')
  self.assertFalse(any(a['id']=='customer' for a in self.s['accounts']))
  self.assertFalse(any(v['id']==voucher['id'] for v in self.s['vouchers']))
 def test_migration_preserves_all_sources_and_is_idempotent(self):
  old=server.demo_state();before=copy.deepcopy(old);ledger.migrate(old)
  for k in before:self.assertEqual(old[k],before[k])
  snapshot=copy.deepcopy(old);ledger.migrate(old);self.assertEqual(old,snapshot)
  self.assertEqual(len(old['vouchers']),len(old['invoices'])+len(old['expenses']))
  self.assertEqual(old['stockMoves'],[]);ledger.validate(old)
  for i in before['invoices']:
   v=ledger.get(old,'vouchers','legacy-sale-'+i['id'])
   expected=i['net']+i['vat']-sum(a['amount'] for a in before['allocations'] if a['invoiceId']==i['id'])
   self.assertEqual(ledger.balance(old,v),max(0,expected))
 def test_xlsx_is_branded_numeric_and_does_not_leak_internal_cost(self):
  pid=self.product();r,_=self.run('quote.save',accountId='customer',date='2026-10-01',title='=BAD()',items=[dict(name='철제함',productId=pid,quantity=2,unitPrice=150000)])
  result,data=self.run('erp.document',kind='quote',id=r['recordId'],format='xlsx')
  ws=load_workbook(io.BytesIO(data)).active
  self.assertEqual(ws['G14'].value,300000);self.assertEqual(ws['H14'].value,30000)
  self.assertEqual(ws['E14'].value,'EA')
  values=' '.join(str(c.value) for row in ws for c in row if c.value is not None)
  self.assertNotIn('100000',values);self.assertNotIn('200000',values)
  self.assertEqual(ws['B7'].value,'=BAD()');self.assertEqual(ws['B7'].data_type,'s')
  again,no_content=self.run('erp.document',kind='quote',id=r['recordId'],format='xlsx');self.assertEqual(result,again);self.assertIsNone(no_content)
 def test_confirmed_edit_cross_customer_and_payment_overage_blocked(self):
  v=self.voucher('sale');self.run('erp.voucher.confirm',id=v['id'])
  with self.assertRaises(ValueError):self.run('erp.voucher.save',id=v['id'])
  journal=copy.deepcopy(self.s['journalEntries'])
  self.run('erp.voucher.info',id=v['id'],title='수정한 건명',attention='새 참조',deliveryTerms='협의',paymentTerms='계좌 이체',notes='표시 정보 수정')
  v=ledger.get(self.s,'vouchers',v['id']);self.assertEqual(v['title'],'수정한 건명');self.assertEqual(v['attention'],'새 참조')
  self.assertEqual(self.s['journalEntries'],journal)
  with self.assertRaises(ValueError):self.run('erp.payment.save',direction='in',accountId='supplier',date='2026-10-01',amount=10,voucherId=v['id'])
  with self.assertRaises(ValueError):self.run('erp.payment.save',direction='in',accountId='customer',date='2026-10-01',amount=999999,voucherId=v['id'])
 def test_future_dated_sale_with_earlier_payment_has_one_consistent_balance(self):
  made,_=self.run('erp.voucher.save',kind='sale',accountId='customer',date='2026-10-13',dueDate='2026-10-13',title='선수금 완료 거래',items=[dict(name='제어반',quantity=1,unitPrice=2353000,vatRate=0)])
  v=ledger.get(self.s,'vouchers',made['recordId']);self.run('erp.voucher.confirm',id=v['id'],date='2026-10-13')
  self.run('erp.payment.save',direction='in',accountId='customer',date='2026-10-11',amount=v['amount'],voucherId=v['id'])
  self.assertEqual(ledger.balance(self.s,v),0)
  self.assertEqual(ledger.balance(self.s,v,'2026-10-11'),v['amount'])
  before=copy.deepcopy(self.s)
  with self.assertRaisesRegex(ValueError,'실제 남은 금액은 0원'):
   self.run('erp.payment.save',direction='in',accountId='customer',date='2026-10-11',amount=5,voucherId=v['id'])
  self.assertEqual(self.s,before)
 def test_voucher_delete_removes_generated_documents_and_reopens_order(self):
  q,_=self.run('quote.save',accountId='customer',date='2026-10-01',title='삭제할 건',attention='원참조',items=[dict(name='철제함',quantity=1,unitPrice=150000)])
  o,_=self.run('quote.convert',quoteId=q['recordId'],date='2026-10-01',orderType='DIS')
  made,_=self.run('erp.order.voucher',orderId=o['recordId'],date='2026-10-01');vid=made['recordId']
  self.run('erp.voucher.confirm',id=vid)
  generated={'documentId':'doc-'+'1'*32};self.s['documents'].append(dict(id=generated['documentId'],generated=True,name='거래명세서.pdf',storedName=generated['documentId']+'.pdf',date='2026-10-01',status='등록 완료',voucherId=vid,orderId=o['recordId']))
  self.s['documents'].append(dict(id='uploaded-delete-test',name='첨부.pdf',date='2026-10-01',status='등록 완료',voucherId=vid,orderId=o['recordId']))
  result,_=self.run('erp.voucher.delete',id=vid)
  self.assertEqual(result['deletedDocumentIds'],[generated['documentId']])
  self.assertFalse(any(v['id']==vid for v in self.s['vouchers']))
  self.assertFalse(any(d['id']==generated['documentId'] for d in self.s['documents']))
  attached=next(d for d in self.s['documents'] if d['id']=='uploaded-delete-test')
  self.assertEqual(attached['voucherId'],'');self.assertEqual(attached['status'],'전표 삭제됨 · 파일 보존')
  self.assertFalse(any(j['sourceId'] in (vid,'reverse-'+vid) for j in self.s['journalEntries']))
  self.assertEqual(workflow.get(self.s,'orders',o['recordId'])['status'],'주문 접수')
  replacement,_=self.run('erp.order.voucher',orderId=o['recordId'],date='2026-10-01')
  self.assertNotEqual(replacement['recordId'],vid)
 def test_voucher_delete_cascades_active_fulfillment_and_payment_history(self):
  pid=self.product();fulfilled=self.voucher('sale',pid=pid);made,_=self.fulfill(fulfilled,1)
  self.run('erp.voucher.delete',id=fulfilled['id'])
  self.assertFalse(any(f.get('voucherId')==fulfilled['id'] for f in self.s['fulfillments']))
  self.assertFalse(any(m.get('sourceId') in (made['fulfillmentId'],'reverse-'+made['fulfillmentId']) for m in self.s['stockMoves']))
  paid=self.voucher('sale');self.run('erp.voucher.confirm',id=paid['id'])
  payment,_=self.run('erp.payment.save',direction='in',accountId='customer',date='2026-10-01',amount=1,voucherId=paid['id'])
  self.run('erp.voucher.delete',id=paid['id'])
  self.assertFalse(any(p['id']==payment['recordId'] for p in self.s['payments']))
  self.assertFalse(any(a.get('voucherId')==paid['id'] for p in self.s['payments'] for a in p.get('allocations',[])))
 def test_order_delete_cascades_voucher_fulfillment_payment_stock_and_documents(self):
  pid=self.product();self.run('erp.stock.opening',productId=pid,date='2026-09-01',quantity=3,value=300000)
  q,_=self.run('quote.save',accountId='customer',date='2026-10-01',title='일괄 삭제 주문',items=[dict(name='철제함',productId=pid,quantity=1,unitPrice=150000)])
  made_order,_=self.run('quote.convert',quoteId=q['recordId'],date='2026-10-01',orderType='DIS');oid=made_order['recordId']
  made_voucher,_=self.run('erp.order.voucher',orderId=oid,date='2026-10-01');vid=made_voucher['recordId']
  self.run('erp.voucher.confirm',id=vid);fulfilled,_=self.run('erp.fulfill',id=vid,date='2026-10-01',items=[dict(lineId=ledger.get(self.s,'vouchers',vid)['items'][0]['id'],quantity=1)])
  paid,_=self.run('erp.payment.save',direction='in',accountId='customer',date='2026-10-01',amount=165000,voucherId=vid)
  doc_id='doc-order-cascade';self.s['documents'].append(dict(id=doc_id,generated=True,name='거래명세서.pdf',storedName=doc_id+'.pdf',date='2026-10-01',status='등록 완료',voucherId=vid,orderId=oid))
  self.s['taxDrafts'].append(dict(id='td-unissued',orderId=oid,status='draft',mode='',stateCode=None,payload=None,approval='',providerApproval='',taxRecordId=''))
  self.s['taxDrafts'].append(dict(id='td-simulation',orderId=oid,status='issued',mode='simulation',stateCode=304,payload={'mock':True},approval='SIM-001'))
  self.s['taxRecords'].append(dict(id='tr-sandbox',orderId=oid,mode='sandbox',approval='SANDBOX-001',providerApproval=''))
  self.s['invoices'].append(dict(id='inv-sandbox',orderId=oid,taxRecordId='tr-sandbox'))
  ledger.get(self.s,'vouchers',vid)['approvalNumber']='SANDBOX-001'
  self.assertEqual(ledger.stock(self.s,pid)['quantity'],2)
  result,_=self.run('order.delete',id=oid)
  self.assertEqual(result['deleted'],dict(vouchers=1,fulfillments=1,payments=1,taxLinks=3))
  self.assertEqual(result['deletedDocumentIds'],[doc_id])
  self.assertFalse(any(o['id']==oid for o in self.s['orders']))
  self.assertFalse(any(v['id']==vid for v in self.s['vouchers']))
  self.assertFalse(any(f['id']==fulfilled['fulfillmentId'] for f in self.s['fulfillments']))
  self.assertFalse(any(p['id']==paid['recordId'] for p in self.s['payments']))
  self.assertFalse(any(d['id']==doc_id for d in self.s['documents']))
  self.assertEqual(ledger.stock(self.s,pid)['quantity'],3)
  self.assertTrue(any(record['id']==q['recordId'] for record in self.s['quotes']))
  self.assertFalse(any(d['id'] in ('td-unissued','td-simulation') for d in self.s['taxDrafts']))
  self.assertFalse(any(r['id']=='tr-sandbox' for r in self.s['taxRecords']))
 def test_order_delete_only_blocks_real_or_uncertain_tax_transmission(self):
  q,_=self.run('quote.save',accountId='customer',date='2026-10-01',title='발급 확인 주문',items=[dict(name='품목',quantity=1,unitPrice=100)])
  made_order,_=self.run('quote.convert',quoteId=q['recordId'],date='2026-10-01',orderType='DIS');oid=made_order['recordId']
  self.s['taxDrafts'].append(dict(id='td-uncertain',orderId=oid,status='unknown',mode='production',stateCode=None,payload={'key':'pending'},approval=''))
  before=copy.deepcopy(self.s)
  with self.assertRaisesRegex(ValueError,'해당 메뉴에서 전송 상태를 확인'):
   self.run('order.delete',id=oid)
  self.assertEqual(self.s,before)
 def test_order_delete_blocks_a_payment_shared_with_another_transaction(self):
  q,_=self.run('quote.save',accountId='customer',date='2026-10-01',title='공유 결제 주문',items=[dict(name='품목',quantity=1,unitPrice=100)])
  made_order,_=self.run('quote.convert',quoteId=q['recordId'],date='2026-10-01',orderType='DIS');oid=made_order['recordId']
  made_voucher,_=self.run('erp.order.voucher',orderId=oid,date='2026-10-01');linked=ledger.get(self.s,'vouchers',made_voucher['recordId']);self.run('erp.voucher.confirm',id=linked['id'])
  other=self.voucher('sale',price=100);self.run('erp.voucher.confirm',id=other['id'])
  payment,_=self.run('erp.payment.save',direction='in',accountId='customer',date='2026-10-01',amount=2)
  self.run('erp.payment.allocate',id=payment['recordId'],voucherId=linked['id'],amount=1,date='2026-10-01')
  self.run('erp.payment.allocate',id=payment['recordId'],voucherId=other['id'],amount=1,date='2026-10-01')
  before=copy.deepcopy(self.s)
  with self.assertRaisesRegex(ValueError,'다른 거래에도 나누어 연결한'):self.run('order.delete',id=oid)
  self.assertEqual(self.s,before)
 def test_personal_draw_is_not_business_cost(self):
  r,_=self.run('erp.voucher.save',kind='withdrawal',date='2026-10-01',title='생활비',items=[dict(name='생활비',quantity=1,unitPrice=100000,vatRate=0)])
  self.run('erp.voucher.confirm',id=r['recordId'])
  self.assertFalse(any(r['account']=='expense' for j in self.s['journalEntries'] for r in j['rows']))
 def test_voucher_and_payment_cancellation_reverse_journal_and_reopen_balance(self):
  v=self.voucher('expense');self.run('erp.voucher.confirm',id=v['id'])
  r,_=self.run('erp.payment.save',direction='out',accountId='supplier',date='2026-10-01',amount=v['amount'],voucherId=v['id'])
  self.run('erp.payment.void',id=r['recordId'],date='2026-10-02',reason='오입력')
  self.assertEqual(ledger.balance(self.s,v),v['amount'])
  self.run('erp.voucher.void',id=v['id'],date='2026-10-02',reason='오입력')
  net={a:sum(r['debit']-r['credit'] for j in self.s['journalEntries'] for r in j['rows'] if r['account']==a) for a in ledger.ACCOUNTS}
  self.assertTrue(all(n==0 for n in net.values()))

 def test_supply_before_confirmation_can_be_cancelled_without_stranded_balances(self):
  pid=self.product();v=self.voucher('purchase',3,100000,pid)
  r,_=self.fulfill(v,3);self.run('erp.voucher.confirm',id=v['id'])
  self.run('erp.fulfillment.void',id=r['fulfillmentId'],date='2026-10-02',reason='입고 정정')
  self.assertEqual(ledger.stock(self.s,pid),dict(quantity=0,value=0))
  self.run('erp.voucher.void',id=v['id'],date='2026-10-02',reason='거래 취소')
  for a in ledger.ACCOUNTS:
   self.assertEqual(sum(r['debit']-r['credit'] for j in self.s['journalEntries'] for r in j['rows'] if r['account']==a),0,a)
 def test_moving_average_and_reverse_latest_receipts(self):
  pid=self.product();v1=self.voucher('purchase',1,100000,pid);v2=self.voucher('purchase',1,200000,pid)
  r1,_=self.fulfill(v1,1);r2,_=self.fulfill(v2,1)
  with self.assertRaises(ValueError):self.run('erp.fulfillment.void',id=r1['fulfillmentId'],date='2026-10-02',reason='순서 오류')
  self.run('erp.fulfillment.void',id=r2['fulfillmentId'],date='2026-10-02',reason='취소')
  self.run('erp.fulfillment.void',id=r1['fulfillmentId'],date='2026-10-02',reason='취소')
  self.assertEqual(ledger.stock(self.s,pid),dict(quantity=0,value=0))
 def test_confirmation_can_record_earlier_invoice_after_recorded_supply(self):
  v=self.voucher('sale');self.fulfill(v,1,'2026-10-03')
  self.run('erp.voucher.confirm',id=v['id'],date='2026-10-01')
  def balance(a,dt):return sum(r['debit']-r['credit'] for j in self.s['journalEntries'] if j['date']<=dt for r in j['rows'] if r['account']==a)
  self.assertEqual(balance('deferredSales','2026-10-02'),-150000)
  self.assertEqual(balance('revenue','2026-10-02'),0)
  self.assertEqual(balance('deferredSales','2026-10-03'),0)
 def test_cancelled_payment_preserves_historical_balance(self):
  v=self.voucher('sale');self.run('erp.voucher.confirm',id=v['id'])
  p,_=self.run('erp.payment.save',direction='in',accountId='customer',date='2026-10-01',amount=v['amount'],voucherId=v['id'])
  self.run('erp.payment.void',id=p['recordId'],date='2026-10-03',reason='오입금')
  self.assertEqual(ledger.balance(self.s,v,'2026-10-02'),0)
  self.assertEqual(ledger.balance(self.s,v,'2026-10-04'),v['amount'])

 def test_payment_and_fulfillment_can_be_edited_and_deleted(self):
  pid=self.product();self.run('erp.stock.opening',productId=pid,date='2026-09-01',quantity=3,value=300000)
  v=self.voucher('sale',2,150000,pid);self.run('erp.voucher.confirm',id=v['id'])
  payment,_=self.run('erp.payment.save',direction='in',accountId='customer',date='2026-10-01',amount=100000,voucherId=v['id'])
  self.run('erp.payment.save',id=payment['recordId'],direction='in',accountId='customer',date='2026-10-02',amount=120000,memo='수정',voucherId='')
  saved=ledger.get(self.s,'payments',payment['recordId']);self.assertEqual((saved['date'],saved['amount'],saved['memo']),('2026-10-02',120000,'수정'))
  made,_=self.fulfill(v,1)
  updated,_=self.run('erp.fulfillment.update',id=made['fulfillmentId'],date='2026-10-02',items=[dict(lineId=v['items'][0]['id'],quantity=2)])
  self.assertEqual(ledger.get(self.s,'fulfillments',updated['fulfillmentId'])['items'][0]['quantity'],2)
  self.run('erp.fulfillment.delete',id=updated['fulfillmentId']);self.assertEqual(ledger.stock(self.s,pid)['quantity'],3)
  self.run('erp.payment.delete',id=payment['recordId']);self.assertEqual(ledger.balance(self.s,v),v['amount'])

 def test_confirmed_voucher_full_edit_resets_dependencies_to_draft(self):
  pid=self.product();v=self.voucher('sale',1,150000,pid);self.run('erp.voucher.confirm',id=v['id'])
  self.fulfill(v,1);self.run('erp.payment.save',direction='in',accountId='customer',date='2026-10-01',amount=100,voucherId=v['id'])
  self.run('erp.voucher.save',id=v['id'],kind='sale',accountId='customer',date='2026-10-01',dueDate='2026-10-05',confirmationDueDate='2026-10-01',title='전체 수정',items=[dict(name='수정 품목',quantity=2,unitPrice=200000,vatRate=0)])
  edited=ledger.get(self.s,'vouchers',v['id']);self.assertEqual(edited['status'],'draft');self.assertEqual(edited['amount'],400000)
  self.assertFalse(any(f['voucherId']==v['id'] for f in self.s['fulfillments']))
  self.assertFalse(any(a['voucherId']==v['id'] for p in self.s['payments'] for a in p.get('allocations',[])))
 def test_original_upload_can_attach_without_another_purchase(self):
  v=self.voucher('purchase');self.s['documents'].append(dict(id='uploaded',name='공급처 거래명세서.pdf',date='2026-10-01',status='확인 필요'))
  self.run('erp.document.attach',id=v['id'],documentId='uploaded');self.assertEqual(len(self.s['vouchers']),1)
  self.assertEqual(self.s['documents'][-1]['voucherId'],v['id'])

 def test_order_statement_keeps_quote_terms_and_does_not_claim_early_delivery(self):
  r,_=self.run('quote.save',accountId='customer',date='2026-10-01',title='A건',attention='구매 담당자',paymentTerms='착수금 30%',deliveryTerms='협의',items=[dict(name='철제함',quantity=1,unitPrice=150000)])
  o,_=self.run('quote.convert',quoteId=r['recordId'],date='2026-10-01',orderType='DIS')
  v,_=self.run('erp.order.voucher',orderId=o['recordId'],date='2026-10-01')
  from erp_documents import snapshot
  doc=snapshot(self.s,'voucher',v['recordId'])
  self.assertEqual(doc['attention'],'구매 담당자');self.assertEqual(doc['paymentTerms'],'착수금 30%')
  self.assertEqual(doc['documentDateLabel'],'작성일');self.assertTrue(doc['orderNumber'].startswith('DIS-'))

 def test_initialization_needs_explicit_confirmation_and_preserves_state_on_refusal(self):
  state=server.demo_state();workflow.normalize(state);before=copy.deepcopy(state)
  with self.assertRaises(ValueError):ledger.apply(state,'erp.initialize',{},'demo')
  self.assertEqual(state,before)



 def test_manual_task_status_roundtrip_and_legacy_compatibility(self):
  fields=dict(accountId='customer',title='상태 검증',date='2026-10-05')
  result,_=self.run('erp.task.save',**fields);tid=result['recordId']
  self.assertEqual(ledger.get(self.s,'tasks',tid)['status'],'not_started')
  for status in ('in_progress','completed','not_started'):
   self.run('erp.task.save',id=tid,status=status,**fields)
   task=ledger.get(self.s,'tasks',tid)
   self.assertEqual(task['status'],status);self.assertEqual(task['done'],status=='completed')
  self.run('erp.task.save',id=tid,**fields)
  self.assertEqual(ledger.get(self.s,'tasks',tid)['status'],'not_started')
  for done in (True,False):
   self.run('erp.task.save',id=tid,done=done,**fields)
   self.assertEqual(ledger.get(self.s,'tasks',tid)['status'],'completed' if done else 'in_progress')
  before=copy.deepcopy(self.s)
  with self.assertRaises(ValueError):self.run('erp.task.save',id=tid,status='invalid',**fields)
  self.assertEqual(self.s,before)

if __name__=='__main__':unittest.main()
