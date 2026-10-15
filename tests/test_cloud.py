"""Cloud integration checks run against a dedicated Postgres database/namespace."""
import base64
import io
import html
import json
import os
import secrets
import tempfile
import time
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from urllib.parse import parse_qs, urlsplit
from unittest.mock import patch

import psycopg
from openpyxl import Workbook
import cloud
from scripts import cloud_admin


class CloudBoundaryTests(unittest.TestCase):
    def setUp(self):
        self.client = cloud.app.test_client()
        self.base = 'https://growth.example'

    def test_no_anonymous_access_to_records_files_or_source(self):
        with patch('cloud.db', side_effect=AssertionError('must not access DB')):
            for path in ('/api/state', '/api/config', '/api/backup-manifest', '/api/documents/x/file', '/api/templates/quote.xlsx', '/api/templates/statement.pdf'):
                self.assertEqual(self.client.get(path, base_url=self.base).status_code, 401)
            self.assertEqual(self.client.get('/tax.js',base_url=self.base).status_code,303)
            self.assertEqual(self.client.get('/', base_url=self.base).status_code, 303)

    def test_authenticated_lookup_module_is_served_without_exposing_source(self):
        self.client.set_cookie(cloud.COOKIE, 'test-lookup-session', domain='growth.example')
        with patch('cloud.db') as connection:
            connection.return_value.__enter__.return_value.execute.return_value.fetchone.return_value=('admin','관리자')
            response=self.client.get('/lookup.js',base_url=self.base)
            self.assertEqual(response.status_code,200)
            self.assertIn('createLookup',response.text)
            draft_response=self.client.get('/drafts.js',base_url=self.base)
            self.assertEqual(draft_response.status_code,200)
            self.assertIn('createDrafts',draft_response.text)
            self.assertEqual(self.client.get('/ledger.py',base_url=self.base).status_code,404)

    def test_health_and_login_work_before_database_provisioning(self):
        self.assertEqual(self.client.get('/api/health', base_url=self.base).status_code, 200)
        self.assertEqual(self.client.get('/login', base_url=self.base).status_code, 200)

    def test_state_reads_after_catalog_deletion_preserve_saved_state(self):
        import ledger
        state = cloud.core.empty_state()
        ledger.migrate(state)
        cloud.merge_products(state)
        ledger.apply(state, 'erp.product.delete', {'id': state['products'][0]['id']}, 'live')
        raw = json.dumps(state)
        for scope in ('learning:demo', 'business:live'):
            with self.subTest(scope=scope), patch('cloud.db') as connection, patch('cloud.state_row', return_value=(42, raw)):
                first = cloud.read_state(scope)
                self.assertEqual(first, {'revision': 42, 'state': cloud.workflow.normalize(json.loads(raw))})
                self.assertEqual(cloud.read_state(scope), first)
                connection.return_value.__enter__.return_value.execute.assert_not_called()
                cloud.core.validate_state(first['state'])

    def test_cross_origin_and_missing_origin_writes_denied(self):
        for headers in ({}, {'Origin': 'https://attacker.example'}):
            response = self.client.post('/api/auth/login', json={}, headers=headers, base_url=self.base)
            self.assertEqual(response.status_code, 403)

    def test_deployment_selects_its_own_workspace_and_namespace(self):
        with patch.dict(os.environ, {}, clear=True):
            with cloud.app.test_request_context('/api/state'):
                self.assertEqual(cloud.scope(), 'learning:demo')
            with patch.dict(os.environ, {'GROWTH_ENABLE_LIVE': '1'}):
                with cloud.app.test_request_context('/api/state'):
                    self.assertEqual(cloud.scope(), 'business:live')
                with cloud.app.test_request_context('/api/state?workspace=demo'):
                    self.assertIsNone(cloud.scope())

    def test_storage_failure_has_no_connection_details(self):
        self.client.set_cookie(cloud.COOKIE, 'a-test-session', domain='growth.example')
        with patch('cloud.db', side_effect=RuntimeError('postgresql://secret:password@private-db')):
            response = self.client.get('/api/state', base_url=self.base)
        self.assertEqual(response.status_code, 503)
        self.assertNotIn('password', response.text)


@unittest.skipUnless(os.environ.get('GROWTH_TEST_DATABASE_URL'), 'Dedicated Postgres connection required')
class CloudDatabaseTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.prefix = 'test-' + secrets.token_hex(8)
        cls.user = cls.prefix + '-owner'
        cls.env = patch.dict(os.environ, {'GROWTH_DATA_NAMESPACE': cls.prefix, 'GROWTH_ENABLE_LIVE': '1', 'GROWTH_FILE_STORAGE': 'postgres'})
        cls.env.start()
        cls.connection = patch('cloud.db', lambda: psycopg.connect(os.environ['GROWTH_TEST_DATABASE_URL'], connect_timeout=10))
        cls.connection.start()
        cloud.setup_schema()

    @classmethod
    def tearDownClass(cls):
        with cloud.db() as con:
            con.execute('DELETE FROM growth_sessions WHERE username LIKE %s', (cls.prefix + '%',))
            con.execute('DELETE FROM growth_users WHERE username LIKE %s', (cls.prefix + '%',))
            con.execute('DELETE FROM growth_files WHERE scope LIKE %s', (cls.prefix + '%',))
            con.execute('DELETE FROM growth_workspaces WHERE scope LIKE %s', (cls.prefix + '%',))
            con.execute('DELETE FROM growth_limits WHERE key IN (%s,%s)', ('user:' + cloud.sha(cls.user), 'ip:' + cloud.sha('127.0.0.1')))
        cls.connection.stop()
        cls.env.stop()

    def setUp(self):
        self.client = cloud.app.test_client()
        self.base = 'https://growth.example'
        self.headers = {'Origin': self.base}
        self.password = 'test-only-long-password-47'
        with cloud.db() as con:
            con.execute('DELETE FROM growth_limits WHERE key IN (%s,%s)', ('user:' + cloud.sha(self.user), 'ip:' + cloud.sha('127.0.0.1')))
            con.execute('INSERT INTO growth_users (username,display_name,password_hash,setup_hash,setup_expires) VALUES (%s,%s,%s,NULL,NULL) ON CONFLICT(username) DO UPDATE SET password_hash=EXCLUDED.password_hash,setup_hash=NULL', (self.user, '테스트 계정', cloud.generate_password_hash(self.password)))
        response = self.post('/api/auth/login', {'username': self.user, 'password': self.password})
        self.assertEqual(response.status_code, 200)
        self.assertIn('Secure', response.headers['Set-Cookie'])
        self.assertIn('HttpOnly', response.headers['Set-Cookie'])
        self.assertIn('SameSite=Strict', response.headers['Set-Cookie'])

    def post(self, path, body):
        return self.client.post(path, json=body, base_url=self.base, headers=self.headers)

    def get(self, path):
        return self.client.get(path, base_url=self.base)

    def test_authenticated_blank_template_downloads_do_not_create_ledger_documents(self):
        before=self.get('/api/state').get_json()
        for kind in ('quote','statement'):
            for fmt in ('pdf','xlsx'):
                response=self.get(f'/api/templates/{kind}.{fmt}')
                self.assertEqual(response.status_code,200)
                self.assertIn('attachment',response.headers['Content-Disposition'])
                self.assertTrue(response.data.startswith(b'%PDF' if fmt=='pdf' else b'PK'))
                response.close()
        self.assertEqual(self.get('/api/state').get_json(),before)
        self.assertEqual(self.get('/api/templates/quote.txt').status_code,400)

    def test_state_is_persistent_and_stale_save_cannot_overwrite(self):
        record = self.get('/api/state?workspace=live').json
        record['state']['settings']['name'] = 'Cloud test'
        self.assertEqual(self.post('/api/state?workspace=live', record).status_code, 200)
        self.assertEqual(self.post('/api/state?workspace=live', record).status_code, 409)
        self.assertEqual(self.get('/api/state?workspace=live').json['state']['settings']['name'], 'Cloud test')

    def test_document_upload_atomic_dedup_and_workspace_isolation(self):
        wb = Workbook(); ws = wb.active
        ws['G3'] = 'QA-001'; ws['G4'] = '2026-10-03'; ws['G5'] = '테스트 고객'
        ws['C14'] = '테스트 부품'; ws['AE14'] = 2; ws['AJ14'] = 100; ws['AQ14'] = 200; ws['AZ14'] = 20
        data = io.BytesIO(); wb.save(data)
        body = {'name': '테스트.xlsx', 'content': base64.b64encode(data.getvalue()).decode()}
        first = self.post('/api/upload?workspace=live', body)
        self.assertEqual(first.status_code, 200, first.text)
        duplicate = self.post('/api/upload?workspace=live', body)
        self.assertTrue(duplicate.json['duplicate'])
        self.assertEqual(first.json['document']['id'], duplicate.json['document']['id'])
        ident = first.json['document']['id']
        self.assertEqual(self.get('/api/documents/' + ident + '/file?workspace=live').data, data.getvalue())
        self.assertEqual(self.get('/api/documents/' + ident + '/file?workspace=demo').status_code, 403)
        snapshot = self.get('/api/backup-manifest?workspace=live').json
        self.assertTrue(any(d['id'] == ident for d in snapshot['state']['documents']))
        snapshot['state']['documents'][0]['storedName'] = '../../other.pdf'
        self.assertEqual(self.post('/api/state?workspace=live', snapshot).status_code, 400)

    def test_supabase_upload_keeps_only_a_private_pointer_and_deduplicates(self):
        wb = Workbook(); wb.active['G3'] = 'SUPABASE-TEST'
        data = io.BytesIO(); wb.save(data); content = data.getvalue()
        body = {'name': 'storage-test.xlsx', 'content': base64.b64encode(content).decode()}
        objects = {}
        def put(scope, ident, extension, data):
            key = cloud.private_storage.object_path(scope, ident, extension)
            objects[key] = data
            return key
        def get(scope, ident, path, digest):
            cloud.private_storage.verify_path(scope, ident, path)
            self.assertEqual(cloud.hashlib.sha256(objects[path]).hexdigest(), digest)
            return objects[path]
        with patch.dict(os.environ, {'GROWTH_FILE_STORAGE': 'supabase'}), \
                patch.object(cloud.private_storage, 'upload', side_effect=put) as upload, \
                patch.object(cloud.private_storage, 'download', side_effect=get):
            first = self.post('/api/upload?workspace=live', body)
            self.assertEqual(first.status_code, 200, first.text)
            ident = first.json['document']['id']
            duplicate = self.post('/api/upload?workspace=live', body)
            self.assertTrue(duplicate.json['duplicate'])
            self.assertEqual(upload.call_count, 1)
            self.assertEqual(self.get('/api/documents/' + ident + '/file?workspace=live').data, content)
            with cloud.db() as con:
                row = con.execute('SELECT body,storage_provider,storage_path FROM growth_files WHERE scope=%s AND id=%s', (self.prefix + ':live', ident)).fetchone()
            self.assertEqual(bytes(row[0]), b'')
            self.assertEqual(row[1], 'supabase')
            self.assertIn(row[2], objects)
            self.assertNotIn('storage_path', first.text)
            self.assertNotIn('supabase.co', first.text)

    def test_storage_failure_does_not_create_a_document_record(self):
        wb = Workbook(); wb.active['G3'] = 'FAILED-STORE'
        data = io.BytesIO(); wb.save(data)
        body = {'name': 'failure-test.xlsx', 'content': base64.b64encode(data.getvalue()).decode()}
        before = self.get('/api/state?workspace=live').json
        with patch.dict(os.environ, {'GROWTH_FILE_STORAGE': 'supabase'}), \
                patch.object(cloud.private_storage, 'upload', side_effect=RuntimeError('storage unavailable')):
            response = self.post('/api/upload?workspace=live', body)
        self.assertEqual(response.status_code, 503)
        self.assertEqual(self.get('/api/state?workspace=live').json, before)

    def test_live_deployment_starts_empty_and_rejects_demo_access(self):
        with patch.dict(os.environ, {'GROWTH_DATA_NAMESPACE': self.prefix + '-fresh'}):
            self.assertEqual(self.get('/api/config').json['workspace'], 'live')
            state = self.get('/api/state').json['state']
            for key in cloud.core.KINDS:
                self.assertEqual(state[key], [], key)
            self.assertEqual(state['settings']['taxReserves'], {})
            self.assertEqual(self.get('/api/state?workspace=demo').status_code, 403)
            self.assertEqual(self.post('/api/state?workspace=demo', {}).status_code, 403)
            self.assertEqual(self.get('/api/backup-manifest').json['workspace'], 'live')

    def test_live_workspace_disabled_on_learning_deployment(self):
        with patch.dict(os.environ, {'GROWTH_ENABLE_LIVE': '0'}):
            self.assertEqual(self.get('/api/state?workspace=live').status_code, 403)
            self.assertEqual(self.post('/api/upload?workspace=live', {}).status_code, 403)
            self.assertEqual(self.get('/api/state?workspace=demo').status_code, 200)

    def test_logout_revokes_server_session(self):
        old_cookie = self.client.get_cookie(cloud.COOKIE, domain='growth.example').value
        self.assertEqual(self.post('/api/auth/logout', {}).status_code, 200)
        self.client.set_cookie(cloud.COOKIE, old_cookie, domain='growth.example')
        self.assertEqual(self.get('/api/state').status_code, 401)

    def test_invite_is_one_use_and_new_password_works(self):
        token = secrets.token_urlsafe(32)
        with cloud.db() as con:
            con.execute('UPDATE growth_users SET setup_hash=%s,setup_expires=%s WHERE username=%s', (cloud.sha(token), int(time.time()) + 300, self.user))
        body = {'username': self.user, 'token': token, 'password': 'new-long-test-password-47'}
        self.assertEqual(self.post('/api/auth/activate', body).status_code, 200)
        self.assertEqual(self.post('/api/auth/activate', body).status_code, 400)
        self.assertEqual(self.post('/api/auth/login', {'username': self.user, 'password': self.password}).status_code, 401)
        self.assertEqual(self.post('/api/auth/login', body).status_code, 200)

    def test_admin_invite_stores_email_and_keeps_setup_token_in_private_file(self):
        username = self.prefix + '-invite'
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / 'setup.html'
            args = ['cloud_admin.py', '--base-url', self.base, '--username', username,
                    '--display-name', 'Invite test', '--email', 'invite@example.com', '--output', str(output)]
            captured = io.StringIO()
            with patch('sys.argv', args), redirect_stdout(captured):
                cloud_admin.main()
            page = output.read_text()
            self.assertEqual(output.stat().st_mode & 0o777, 0o600)
            self.assertIn('invite@example.com', page)
            link = html.unescape(page.split('href="', 1)[1].split('"', 1)[0])
            token = parse_qs(urlsplit(link).fragment)['setup'][0]
            self.assertNotIn(token, captured.getvalue())
            with cloud.db() as con:
                email, password, digest = con.execute('SELECT email,password_hash,setup_hash FROM growth_users WHERE username=%s', (username,)).fetchone()
            self.assertEqual(email, 'invite@example.com')
            self.assertIsNone(password)
            self.assertEqual(digest, cloud.sha(token))
            body = {'username': username, 'token': token, 'password': self.password}
            self.assertEqual(self.post('/api/auth/activate', body).status_code, 200)
            self.assertEqual(self.get('/api/config').json['user']['username'], username)
            self.assertEqual(self.post('/api/auth/activate', body).status_code, 400)
        with cloud.db() as con:
            con.execute('DELETE FROM growth_limits WHERE key=%s', ('user:' + cloud.sha(username),))

    def test_wrong_password_and_repeated_login_are_rejected(self):
        for _ in range(9):
            self.assertEqual(self.post('/api/auth/login', {'username': self.user, 'password': 'wrong'}).status_code, 401)
        self.assertEqual(self.post('/api/auth/login', {'username': self.user, 'password': self.password}).status_code, 429)

    def test_v2_atomic_stock_commands_and_private_xlsx_export(self):
        with patch.dict(os.environ, {'GROWTH_DATA_NAMESPACE': self.prefix + '-v2'}):
            def command(action, **data):
                revision = self.get('/api/state').json['revision']
                return self.post('/api/workflow', dict(revision=revision, action=action, data=data))
            before = self.get('/api/state').json
            self.assertEqual(command('erp.initialize').status_code, 400)
            self.assertEqual(self.get('/api/state').json, before)
            self.assertEqual(command('erp.initialize', confirmed=True).status_code, 200)
            a = command('erp.account.save', name='통합 장부 테스트').json['recordId']
            p = command('erp.product.save', code='BOX-QA', name='철제함', unit='EA', salePrice=150000, purchasePrice=100000).json['recordId']
            self.assertEqual(command('erp.stock.opening', productId=p, date='2026-10-01', quantity=1, value=100000).status_code, 200)
            v = command('erp.voucher.save', kind='sale', accountId=a, title='원자적 출고 검증', date='2026-10-01', items=[dict(name='철제함', productId=p, quantity=1, unitPrice=150000), dict(name='철제함', productId=p, quantity=1, unitPrice=150000)]).json
            ident = v['recordId']; lines = v['state']['vouchers'][-1]['items']
            self.assertEqual(command('erp.voucher.confirm', id=ident).status_code, 200)
            before = self.get('/api/state').json
            failed = command('erp.fulfill', id=ident, date='2026-10-01', items=[dict(lineId=r['id'], quantity=i+1) for i,r in enumerate(lines)])
            self.assertEqual(failed.status_code, 400, failed.text)
            self.assertEqual(self.get('/api/state').json, before)  # No partial stock or journal write.
            completed = command('erp.fulfill', id=ident, date='2026-10-01', items=[dict(lineId=r['id'], quantity=1) for r in lines])
            self.assertEqual(completed.status_code, 200, completed.text)
            persisted = self.get('/api/state').json['state']
            self.assertEqual(sum(m['quantity'] for m in persisted['stockMoves'] if m['productId']==p), -1)
            self.assertTrue(persisted['fulfillments'][-1]['costProvisional'])
            self.assertEqual(persisted, completed.json['state'])
            before = self.get('/api/state').json
            forged = json.loads(json.dumps(before)); forged['state']['vouchers'][0]['net'] = 1
            self.assertEqual(self.post('/api/state', forged).status_code, 400)
            objects = {}
            def upload(scope, key, extension, data):
                self.assertEqual(extension, '.xlsx')
                objects[key] = data
                return cloud.private_storage.object_path(scope, key, extension)
            with patch.dict(os.environ, {'GROWTH_FILE_STORAGE': 'supabase'}), patch.object(cloud.private_storage, 'upload', side_effect=upload), patch.object(cloud.private_storage, 'download', side_effect=lambda scope, key, path, digest: objects[key]):
                out = command('erp.document', kind='voucher', id=ident, format='xlsx')
                self.assertEqual(out.status_code, 200, out.text)
                file = self.get('/api/documents/' + out.json['documentId'] + '/file?download=1')
                self.assertEqual(file.status_code, 200)
                self.assertIn('spreadsheetml', file.content_type)
                self.assertTrue(file.data.startswith(b'PK'))
                self.assertNotIn('storage_path', out.text)
                count = len(out.json['state']['documents'])
                again = command('erp.document', kind='voucher', id=ident, format='xlsx')
                self.assertEqual(again.json['documentId'], out.json['documentId'])
                self.assertEqual(len(again.json['state']['documents']), count)

    def test_authenticated_user_cannot_download_source_or_env(self):
        for path in ('/.env.local', '/cloud.py', '/server.py', '/data/growth.sqlite3'):
            self.assertEqual(self.get(path).status_code, 404)

    def test_password_change_revokes_other_sessions(self):
        old_cookie = self.client.get_cookie(cloud.COOKIE, domain='growth.example').value
        self.assertEqual(self.post('/api/auth/password', {'currentPassword': self.password, 'password': 'another-new-test-password'}).status_code, 200)
        self.assertEqual(self.get('/api/state').status_code, 200)
        self.client.set_cookie(cloud.COOKIE, old_cookie, domain='growth.example')
        self.assertEqual(self.get('/api/state').status_code, 401)

    def test_quote_delivery_pdf_commands_persist_atomically_and_reject_stale_and_tampering(self):
        import copy
        from pypdf import PdfReader
        with patch.dict(os.environ, {'GROWTH_DATA_NAMESPACE': self.prefix + '-workflow'}):
            original = self.get('/api/state').json
            original['state']['accounts'].append(dict(id='test-customer',name='루미나 캠퍼스 테스트',kind='고객사'))
            self.assertEqual(self.post('/api/state',original).status_code,200)
            def command(action, data):
                record=self.get('/api/state').json
                response=self.post('/api/workflow',dict(action=action,data=data,revision=record['revision']))
                self.assertEqual(response.status_code,200,response.json)
                return response.json
            q=command('quote.save',dict(accountId='test-customer',date='2026-10-03',title='통합 테스트',items=[dict(name='철함',quantity=3,unitPrice=150000)]))
            stale=self.post('/api/workflow',dict(action='quote.save',data={},revision=original['revision']))
            self.assertEqual(stale.status_code,409)
            generated=command('document.generate',dict(kind='quote',id=q['recordId']))
            pdf=self.get('/api/documents/'+generated['documentId']+'/file')
            self.assertEqual(pdf.status_code,200)
            self.assertEqual(pdf.mimetype,'application/pdf')
            self.assertIn('루미나 캠퍼스 테스트',PdfReader(io.BytesIO(pdf.data)).pages[0].extract_text())
            download=self.get('/api/documents/'+generated['documentId']+'/file?download=1')
            self.assertIn('attachment',download.headers['Content-Disposition'])
            self.assertEqual(command('document.generate',dict(kind='quote',id=q['recordId']))['documentId'],generated['documentId'])
            unchanged=self.get('/api/state').json
            forged=copy.deepcopy(unchanged);forged['state']['quotes'][-1]['net']=1
            self.assertEqual(self.post('/api/state',forged).status_code,400)
            self.assertEqual(self.get('/api/state').json,unchanged)
            o=command('quote.convert',dict(quoteId=q['recordId'],date='2026-10-03'))
            order=o['state']['orders'][-1]
            d=command('delivery.create',dict(orderId=order['id'],date='2026-10-04',items=[dict(orderLineId=order['items'][0]['id'],quantity=1)]))
            before=self.get('/api/state').json
            with patch.dict(os.environ, {'GROWTH_FILE_STORAGE':'supabase'}), patch('cloud.private_storage.upload',side_effect=cloud.private_storage.StorageUnavailable('unavailable')):
                failed=self.post('/api/workflow',dict(action='document.generate',data=dict(kind='delivery',id=d['recordId']),revision=before['revision']))
            self.assertEqual(failed.status_code,503)
            self.assertEqual(self.get('/api/state').json,before)
            stored={}
            def upload(scope,ident,extension,content):
                stored['content']=content
                return 'files/test/'+ident+'.pdf'
            with patch.dict(os.environ, {'GROWTH_FILE_STORAGE':'supabase'}), patch('cloud.private_storage.upload',side_effect=upload):
                delivery_pdf=command('document.generate',dict(kind='delivery',id=d['recordId']))
            with cloud.db() as con:
                row=con.execute('SELECT body,storage_provider,storage_path FROM growth_files WHERE scope=%s AND id=%s',(self.prefix+'-workflow:live',delivery_pdf['documentId'])).fetchone()
            self.assertEqual(bytes(row[0]),b'');self.assertEqual(row[1],'supabase');self.assertTrue(row[2].endswith('.pdf'))
            with patch('cloud.private_storage.download',return_value=stored['content']):
                self.assertEqual(self.get('/api/documents/'+delivery_pdf['documentId']+'/file').data,stored['content'])
            self.assertIn('165,000',PdfReader(io.BytesIO(stored['content'])).pages[0].extract_text())
            final=self.get('/api/backup-manifest').json['state']
            self.assertEqual(len(final['documents']),2)
            self.assertEqual(final['invoices'],[])
            self.assertEqual(final['deliveries'][0]['items'][0]['quantity'],1)
            anonymous=cloud.app.test_client()
            self.assertEqual(anonymous.get('/api/documents/'+delivery_pdf['documentId']+'/file',base_url=self.base).status_code,401)
            self.assertEqual(anonymous.post('/api/workflow',json={},headers=self.headers,base_url=self.base).status_code,401)


    def test_tax_outbox_is_committed_before_network_and_survives_unknown_result(self):
        import copy
        import tax_provider
        with patch.dict(os.environ, {'GROWTH_DATA_NAMESPACE':self.prefix+'-tax','GROWTH_ENABLE_LIVE':'0','GROWTH_TAX_MODE':'simulation'}):
            current=self.get('/api/state').json
            current['state']['settings'].update(businessNumber='8888800001',representative='테스트')
            current['state']['accounts'].append(dict(id='tax-customer',name='가상 세금 고객',kind='고객사',businessNumber='8888800003',representative='테스트'))
            self.assertEqual(self.post('/api/state',current).status_code,200)
            def command(action,data,external=False):
                r=self.get('/api/state').json
                response=self.post('/api/tax' if external else '/api/workflow',dict(action=action,data=data,revision=r['revision']))
                self.assertEqual(response.status_code,200,response.json)
                return response.json
            order=command('order.save',dict(accountId='tax-customer',title='클라우드 발급 테스트',date='2026-10-03',items=[dict(name='가상 품목',quantity=1,unitPrice=200000)]))
            draft=command('tax.draft',dict(orderId=order['recordId'],date='2026-10-03'))
            def timeout(op):
                # Separate connection sees the committed outbox before any external call.
                with cloud.db() as con:
                    raw=con.execute('SELECT body FROM growth_workspaces WHERE scope=%s',(self.prefix+'-tax:demo',)).fetchone()[0]
                    saved=json.loads(raw)['taxDrafts'][0]
                    self.assertEqual(saved['status'],'submitting')
                    self.assertEqual(saved['payload']['invoicerMgtKey'],op['key'])
                return dict(error='test timeout',uncertain=True)
            with patch('cloud.tax_provider.execute',side_effect=timeout):
                result=command('draft.issue',dict(id=draft['recordId'],confirmed=True),True)
            self.assertEqual(result['state']['taxDrafts'][0]['status'],'unknown')
            revision=result['revision']
            forbidden=self.post('/api/tax',dict(action='draft.issue',data=dict(id=draft['recordId'],confirmed=True),revision=revision))
            self.assertEqual(forbidden.status_code,400)
            status=command('draft.status',dict(id=draft['recordId']),True)
            d=status['state']['taxDrafts'][0]
            self.assertEqual(d['stateCode'],304)
            self.assertEqual(sum(i.get('taxRecordId')==d['taxRecordId'] for i in status['state']['invoices']),1)
            corrected=command('tax.correct',dict(id=d['id'],modifyCode=3,date='2026-10-03',net=-50000,vat=-5000))
            result=command('draft.issue',dict(id=corrected['recordId'],confirmed=True),True)
            invoices=[i for i in result['state']['invoices'] if i['accountId']=='tax-customer']
            self.assertEqual(sum(i['net']+i['vat'] for i in invoices),165000)
            current=self.get('/api/state').json;forged=copy.deepcopy(current);forged['state']['taxDrafts'][0]['net']=1
            self.assertEqual(self.post('/api/state',forged).status_code,400)
            self.assertEqual(self.get('/api/state').json,current)
            stale=self.post('/api/tax',dict(action='draft.status',data=dict(id=d['id']),revision=revision))
            self.assertEqual(stale.status_code,409)
            anon=cloud.app.test_client()
            self.assertEqual(anon.post('/api/tax',json={},base_url=self.base,headers=self.headers).status_code,401)

    def test_statement_sales_manual_tax_and_supplier_pair_use_authenticated_atomic_workflow(self):
        import copy
        with patch.dict(os.environ, {'GROWTH_DATA_NAMESPACE':self.prefix+'-statements','GROWTH_TAX_MODE':'disabled'}), patch('cloud.tax_provider.execute',side_effect=AssertionError('no provider calls')):
            current=self.get('/api/state').json
            current['state']['accounts'].extend([dict(id='customer',name='가상 고객',kind='고객사'),dict(id='vendor',name='가상 공급처',kind='매입처')])
            self.assertEqual(self.post('/api/state',current).status_code,200)
            def command(action, **data):
                current=self.get('/api/state').json
                response=self.post('/api/workflow',dict(action=action,data=data,revision=current['revision']))
                self.assertEqual(response.status_code,200,response.json)
                return response.json
            o=command('order.save',accountId='customer',title='명세서 테스트',date='2026-10-01',items=[dict(name='분전반',quantity=1,unitPrice=500000)])
            order=o['state']['orders'][0]
            d=command('delivery.create',orderId=order['id'],date='2026-10-02',items=[dict(orderLineId=order['items'][0]['id'],quantity=1)])
            generated=command('document.generate',kind='delivery',id=d['recordId'])
            self.assertEqual(generated['state']['invoices'],[])
            sale=command('statement.confirm',id=d['recordId'],confirmed=True)
            invoice=sale['state']['invoices'][0]
            self.assertEqual((invoice['net'],invoice['vat'],invoice['date']),(500000,50000,'2026-10-02'))
            self.assertEqual(len(command('statement.confirm',id=d['recordId'],confirmed=True)['state']['invoices']),1)
            for done in (True, False):
                changed=command('statement.hometax-task',id=invoice['id'],done=done)
                self.assertEqual(self.get('/api/state').json['state']['invoices'][0]['hometaxTask']['done'],done)
                unchanged=copy.deepcopy(changed['state']);unchanged['invoices'][0].pop('hometaxTask');unchanged['invoices'][0].pop('hometaxTaskHistory')
                self.assertEqual(unchanged,sale['state'])
            forged_task=copy.deepcopy(changed);forged_task['state']['invoices'][0]['hometaxTask']['done']=True
            self.assertEqual(self.post('/api/state',forged_task).status_code,400)
            self.assertEqual(self.post('/api/workflow',dict(action='statement.hometax-task',data=dict(id=invoice['id'],done=True),revision=sale['revision'])).status_code,409)
            tax=command('statement.tax-confirm',invoiceIds=[invoice['id']],date='2026-10-03',confirmed=True)
            self.assertEqual(tax['state']['invoices'][0]['date'],'2026-10-02')
            forged=copy.deepcopy(tax);forged['state']['invoices'][0]['net']=1
            self.assertEqual(self.post('/api/state',forged).status_code,400)
            config=self.get('/api/config').json
            self.assertEqual(config['tax']['mode'],'disabled')
            self.assertEqual(self.get('/statements.js').status_code,200)
            anonymous=cloud.app.test_client()
            self.assertEqual(anonymous.get('/statements.js',base_url=self.base).status_code,303)
            # Add distinct test PDFs via the actual upload endpoint, then pair them.
            from reportlab.pdfgen import canvas
            documents=[]
            for title in ('supplier statement','supplier tax invoice'):
                buffer=io.BytesIO();pdf=canvas.Canvas(buffer);pdf.drawString(40,700,title);pdf.save()
                uploaded=self.post('/api/upload',dict(name=title+'.pdf',content=base64.b64encode(buffer.getvalue()).decode()))
                self.assertEqual(uploaded.status_code,200,uploaded.json);documents.append(uploaded.json['document']['id'])
            data=dict(accountId='vendor',date='2026-10-02',kind='거래명세서',category='자재비',net=300000,vat=30000,confirmed=True)
            cost=command('statement.purchase',documentId=documents[0],**data)
            data['kind']='세금계산서'
            linked=command('statement.purchase',documentId=documents[1],existingExpenseId=cost['recordId'],**data)
            self.assertEqual(len(linked['state']['expenses']),1)
            self.assertEqual(sum(doc.get('expenseId')==cost['recordId'] for doc in linked['state']['documents']),2)


if __name__ == '__main__':
    unittest.main()
