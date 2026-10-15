import copy, io, unittest
from openpyxl import load_workbook
from pypdf import PdfReader
import server, workflow, ledger, erp_documents
from company_profile import SG_DEFAULTS, SG_PREVIOUS, initialize_company
from document_templates import export_template
from pdf_documents import render_pdf


class TemplateExportTests(unittest.TestCase):
    def state(self):
        s=workflow.normalize(server.empty_state());ledger.migrate(s)
        return s

    def workbook(self,content):return load_workbook(io.BytesIO(content)).active
    def pdf(self,content):return PdfReader(io.BytesIO(content))
    def text(self,content):return '\n'.join(p.extract_text() for p in self.pdf(content).pages)

    def test_templates_survive_fresh_empty_ledger_and_have_no_transaction_dependency(self):
        for kind in ('quote','statement'):
            for fmt in ('pdf','xlsx'):
                s=self.state();before=copy.deepcopy(s)
                name,content=export_template(kind,fmt,s['settings'])
                self.assertEqual(s,before);self.assertFalse(s['documents']);self.assertTrue(name.endswith('.'+fmt))
                text=self.text(content) if fmt=='pdf' else ' '.join(str(c.value) for row in self.workbook(content) for c in row if c.value is not None)
                for term in ('그로스 일렉트릭','000-00-00000','김데모','가상시 데모구 그로스길 10','제조,도소매','전기콘트롤박스','010-0000-0000','02-000-0000'):
                    self.assertIn(term,text)
                self.assertNotIn('본 문서는 세금계산서가 아닙니다',text);self.assertNotIn('테스트용',text);self.assertNotIn('루미나 캠퍼스',text);self.assertNotIn('부천시',text);self.assertNotIn('납품장소',text);self.assertNotIn('연결 주문',text)
                if fmt=='xlsx':self.assertEqual(len(self.workbook(content)._images),0)
                else:self.assertEqual(len(self.pdf(content).pages[0].images), 0)

    def test_xlsx_header_totals_memo_and_items_have_visible_rules(self):
        _,content=export_template('quote','xlsx');ws=self.workbook(content)
        for address in ('A4','B4','E4','F4','E5','F5','H5','I5','E6','F6','E7','F7','H7','I7','E8','F8','H8','I8','A10','B10','E10','F10','A11','B11','A13','E13','I13','A14','I19','A20','G20','I20','A22','A23'):
            c=ws[address]
            with self.subTest(address=address):
                self.assertIsNotNone(c.border.top.style);self.assertIsNotNone(c.border.bottom.style)
                self.assertIsNotNone(c.border.left.style);self.assertIsNotNone(c.border.right.style)
        self.assertEqual(ws['E13'].value,'단위');self.assertIsNone(ws['G14'].value)
        self.assertTrue(ws.sheet_properties.pageSetUpPr.fitToPage)

    def test_profile_update_once_retains_later_customization_and_other_suppliers(self):
        previous=dict(SG_DEFAULTS,**SG_PREVIOUS)
        initialize_company(previous);self.assertEqual(previous['address'],SG_DEFAULTS['address'])
        previous['address']='사용자가 나중에 변경한 주소';initialize_company(previous)
        self.assertEqual(previous['address'],'사용자가 나중에 변경한 주소')
        other=dict(name='공급처 A',businessNumber='123-45-67890',address='공급처 주소')
        original=copy.deepcopy(other);initialize_company(other);self.assertEqual(other,original)
        _,pdf=export_template('quote','pdf',other)
        self.assertIn('공급처 A',self.text(pdf));self.assertEqual(len(self.pdf(pdf).pages[0].images),0)

    def test_new_generation_updates_template_without_rewriting_stored_document(self):
        s=self.state();s['accounts']=[dict(id='a',name='루미나 캠퍼스')]
        q=workflow.apply(s,dict(action='quote.save',data=dict(accountId='a',date='2026-10-05',title='도면 A',attention='담당 A',items=[dict(name='철함',quantity=2,unit='EA',unitPrice=150000)])),'demo')[0]['recordId']
        record=ledger.get(s,'quotes',q);record['company'].update(SG_PREVIOUS)
        first,pdf=workflow.apply(s,dict(action='document.generate',data=dict(kind='quote',id=q)),'demo')
        saved=copy.deepcopy(s['documents']);source_company=copy.deepcopy(record['company'])
        self.assertIn(SG_DEFAULTS['address'],self.text(pdf));self.assertIn('루미나 캠퍼스',self.text(pdf))
        s['settings']['address']='추후 변경 주소';s['accounts'][0]['name']='거래처 현재 이름 변경'
        second,pdf2=workflow.apply(s,dict(action='document.generate',data=dict(kind='quote',id=q)),'demo')
        self.assertNotEqual(first['documentId'],second['documentId'])
        self.assertEqual(s['documents'][:len(saved)],saved);self.assertEqual(record['company'],source_company)
        self.assertIn('추후 변경 주소',self.text(pdf2));self.assertIn('루미나 캠퍼스',self.text(pdf2));self.assertNotIn('거래처 현재 이름 변경',self.text(pdf2))
        self.assertIn('담당 A',self.text(pdf2));self.assertIn('330,000',self.text(pdf2))

    def test_purchase_supplier_identity_is_not_replaced_with_our_branding(self):
        s=self.state();s['accounts']=[dict(id='vendor',name='외부 공급처',businessNumber='123-45-67890',address='외부 공급처 주소')]
        result,_=ledger.apply(s,'erp.voucher.save',dict(kind='purchase',accountId='vendor',date='2026-10-05',title='자재',items=[dict(name='철함',quantity=1,unitPrice=100000)]),'demo')
        r=erp_documents.snapshot(s,'voucher',result['recordId'])
        self.assertEqual(r['company']['name'],'외부 공급처');self.assertEqual(r['customer']['name'],SG_DEFAULTS['name'])
        self.assertEqual(len(self.workbook(erp_documents.render_xlsx('voucher',r))._images),0)

    def test_quote_and_statement_exports_hide_categories_and_omit_zero_vat_caption(self):
        s=self.state();s['accounts']=[dict(id='a',name='출력 확인 거래처')]
        item=dict(name='품목 A',spec='규격 A',section='표시금지구분',quantity=1,unitPrice=150000,vatRate=0)
        q=workflow.apply(s,dict(action='quote.save',data=dict(accountId='a',date='2026-10-05',title='견적',items=[item])),'demo')[0]['recordId']
        v=ledger.apply(s,'erp.voucher.save',dict(kind='sale',accountId='a',date='2026-10-05',title='거래',items=[item]),'demo')[0]['recordId']
        for kind,rid in (('quote',q),('voucher',v)):
            record=erp_documents.snapshot(s,kind,rid)
            pdf_text=self.text(render_pdf('quote' if kind=='quote' else 'delivery',record))
            ws=self.workbook(erp_documents.render_xlsx(kind,record))
            xlsx_text=' '.join(str(c.value) for row in ws for c in row if c.value is not None)
            for text in (pdf_text,xlsx_text):
                self.assertNotIn('부가세 0원',text)
                self.assertNotIn('부가세 포함',text)
                self.assertIn('십오만원정',text)
                self.assertNotIn('표시금지구분',text)
                self.assertIn('규격 A',text)
            self.assertEqual(ws['C13'].value,'규격')
            self.assertEqual(ws['C14'].value,'규격 A')

    def test_all_new_documents_omit_removed_footer_including_learning_workspace(self):
        from unittest.mock import patch
        s=self.state();s['accounts']=[dict(id='a',name='출력 확인 거래처')]
        q=workflow.apply(s,dict(action='quote.save',data=dict(accountId='a',date='2026-10-05',title='견적',items=[dict(name='품목',quantity=2,unitPrice=150000)])),'demo')[0]['recordId']
        v=ledger.apply(s,'erp.voucher.save',dict(kind='sale',accountId='a',date='2026-10-05',title='거래',items=[dict(name='품목',quantity=2,unitPrice=150000)]),'demo')[0]['recordId']
        for kind,rid in (('quote',q),('voucher',v)):
            record=erp_documents.snapshot(s,kind,rid)
            for demo in (True,False):
                for fmt in ('pdf','xlsx'):
                    content=render_pdf('quote' if kind=='quote' else 'delivery',record,demo) if fmt=='pdf' else erp_documents.render_xlsx(kind,record,demo)
                    text=self.text(content) if fmt=='pdf' else ' '.join(str(c.value) for row in self.workbook(content) for c in row if c.value is not None)
                    for term in ('본 문서는 세금계산서가 아닙니다','테스트용','고객 발송용이 아닙니다'):
                        self.assertNotIn(term,text)
                    self.assertIn('출력 확인 거래처',text)
        # Template-version change must regenerate files rather than returning a cached old footer.
        for fmt in ('pdf','xlsx'):
            payload=dict(kind='quote',id=q,format=fmt)
            with patch('erp_documents.TEMPLATE_VERSION',4):old,_=erp_documents.generate(s,payload,'demo')
            saved=copy.deepcopy(s['documents'])
            new,content=erp_documents.generate(s,payload,'demo')
            self.assertNotEqual(old['documentId'],new['documentId']);self.assertTrue(content)
            self.assertEqual(s['documents'][:len(saved)],saved)
