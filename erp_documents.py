"""Customer-facing XLSX/PDF exports. Internal costs never appear in these files."""
import copy,hashlib,io,json,re
from pathlib import Path
import ledger
import workflow
from document_templates import TEMPLATE_VERSION


def snapshot(s,kind,ident):
 r=copy.deepcopy(ledger.get(s,'quotes' if kind=='quote' else 'vouchers',ident))
 if kind=='quote':
  ledger.need(r.get('workflowVersion')==2,'기존 견적은 품목을 확인해 새 버전으로 저장해 주세요.')
  r['company']=workflow.profile(s['settings']);r.setdefault('customer',workflow.profile(ledger.get(s,'accounts',r['accountId'])))
 else:
  r['documentDateLabel']='작성일'
  order=next((o for o in s['orders'] if o['id']==r.get('orderId')), {})
  r['orderNumber']=order.get('number',r.get('orderId',''))
  ledger.need(r['status']!='void','취소한 전표입니다.')
  ledger.need(r['kind'] in ('sale','purchase'),'매출 또는 매입 거래명세서를 선택해 주세요.')
  company=workflow.profile(s['settings'])
  customer=r.get('customer') or workflow.profile(ledger.get(s,'accounts',r['accountId']))
  if r['kind']=='purchase':
   r['company'],r['customer']=customer,company
   r['notes']='매입 거래명세서 · 내부 보관용\n'+r.get('notes','')
  else:r['company'],r['customer']=company,customer
  if not r['items']:
   r['items']=[dict(id='legacy-total',name=r['title'],spec='',section='',unit='식',quantity=1,unitPrice=r['net'],net=r['net'],vat=r['vat'],vatRate=10 if r['vat'] else 0)]
  r['notes']=('미확정 전표\n' if r['status']=='draft' else '확정 전표\n')+r.get('notes','')
  if r.get('approvalNumber'):r['notes']+='\n국세청 승인번호 메모: '+r['approvalNumber']
 r.pop('taxEvidence',None);r.pop('orderDifferenceReason',None)
 r['items']=[{k:v for k,v in row.items() if k in ('id','section','name','spec','unit','quantity','unitPrice','net','vat','vatRate')} for row in r['items']]
 return r


def render_xlsx(kind,r,demo=False):
 from openpyxl import Workbook
 from openpyxl.styles import Font,Alignment,Border,Side,PatternFill
 from openpyxl.drawing.image import Image
 from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor,AnchorMarker
 from openpyxl.drawing.xdr import XDRPositiveSize2D
 from openpyxl.utils.units import pixels_to_EMU
 from company_profile import use_sg_seal
 from quote_pdf import korean_won,vat_caption
 w=Workbook();ws=w.active;ws.title='견적서' if kind=='quote' else '거래명세서'
 ws.sheet_view.showGridLines=False;blank=r.get('isTemplate',False)
 for col,width in dict(A=9,B=25,C=20,D=9,E=10,F=14,G=16,H=11,I=16).items():ws.column_dimensions[col].width=width
 thin=Side(style='thin',color='444444');strong=Side(style='medium',color='202020')
 grid=Border(left=thin,right=thin,top=thin,bottom=thin)
 def cell(row,col,value='',bold=False,size=10,align='left'):
  c=ws.cell(row,col,value)
  if isinstance(value,str):c.data_type='s'
  c.font=Font(name='맑은 고딕',size=size,bold=bold,color='202020')
  c.alignment=Alignment(horizontal=align,vertical='center',wrap_text=True)
  return c
 def box(row,start,end,value='',bold=False,size=10,align='left',border=True):
  if end>start:ws.merge_cells(start_row=row,start_column=start,end_row=row,end_column=end)
  c=cell(row,start,value,bold,size,align)
  if border:
   for col in range(start,end+1):
    ws.cell(row,col).border=Border(left=thin if col==start else Side(),right=thin if col==end else Side(),top=thin,bottom=thin)
  return c
 box(1,1,9,'견  적  서' if kind=='quote' else '거 래 명 세 서',True,24,'center',False);ws.row_dimensions[1].height=44
 box(2,1,9,'기본 양식' if blank else '',False,9,border=False);ws.row_dimensions[2].height=18
 # Internal order numbers are deliberately omitted from customer documents.
 box(3,1,9,'',False,9,border=False);ws.row_dimensions[3].height=20
 supplier,buyer=r['company'],r['customer']
 reference=r['number']+(f" / v{r.get('version',1)}" if kind=='quote' and not blank else '')
 left=[('견적서 No' if kind=='quote' else '명세서 No',reference),('작성일',r['date']),('상호',buyer.get('name','')),('건명',r.get('title','')),('참조',r.get('attention',''))]
 for row,(label,value) in enumerate(left,4):
  box(row,1,1,label,True,9,'center');box(row,2,4,value)
  ws.row_dimensions[row].height=32 if row!=5 else 44
 box(9,1,4,'아래와 같이 견적합니다.' if kind=='quote' else '아래와 같이 거래 내역을 확인합니다.',False,10,'center');ws.row_dimensions[9].height=29
 for row,label in ((4,'등록번호'),(5,'상호명'),(6,'주소'),(7,'업태'),(8,'전화번호'),(9,'이메일')):box(row,5,5,label,True,9,'center')
 box(4,6,9,supplier.get('businessNumber',''))
 box(5,6,7,supplier.get('name',''),True);box(5,8,8,'성명',True,9,'center');box(5,9,9,supplier.get('representative',''))
 box(6,6,9,supplier.get('address',''),size=9)
 box(7,6,7,supplier.get('businessType',''));box(7,8,8,'종목',True,9,'center');box(7,9,9,supplier.get('businessItem',''),size=9)
 box(8,6,7,supplier.get('phone',''));box(8,8,8,'팩스',True,9,'center');box(8,9,9,supplier.get('fax',''),size=9)
 box(9,6,9,supplier.get('email',''),size=9)
 if use_sg_seal(supplier):
  im=Image(str(Path(__file__).parent/'assets/branding/sg-seal.png'));im.width=40;im.height=43
  im.anchor=OneCellAnchor(_from=AnchorMarker(col=8,row=4,colOff=pixels_to_EMU(61),rowOff=pixels_to_EMU(6)),ext=XDRPositiveSize2D(pixels_to_EMU(40),pixels_to_EMU(43)))
  ws.add_image(im)
 box(10,1,1,'납기',True,9,'center');box(10,2,4,r.get('deliveryTerms',''))
 box(10,5,5,'지불조건',True,9,'center');box(10,6,9,r.get('paymentTerms',''));ws.row_dimensions[10].height=30
 box(11,1,1,'합계금액',True,9,'center')
 total=r['net']+r['vat'];box(11,2,9,'' if blank else '금 '+korean_won(total)+vat_caption(r['vat'])+'   ₩ '+format(total,','),True,11)
 ws.row_dimensions[11].height=35;ws.row_dimensions[12].height=12
 headers=['No','품명','규격','수량','단위','단가','공급가액','부가세','합계']
 ws.row_dimensions[13].height=27
 for col,name in enumerate(headers,1):
  c=cell(13,col,name,True,10,'center');c.fill=PatternFill('solid',fgColor='F3F4F5');c.border=grid
 count=max(6,len(r['items']))
 for index in range(count):
  row=14+index;item=r['items'][index] if index<len(r['items']) else None
  values=([index+1,item['name'],item.get('spec',''),item['quantity'],item.get('unit',''),item['unitPrice'],item['net'],item['vat'],item['net']+item['vat']] if item and not blank else ['']*9)
  for col,value in enumerate(values,1):
   c=cell(row,col,value);c.border=grid
   if isinstance(value,(int,float)):
    c.number_format='#,##0.###;[Red]-#,##0.###' if col==4 and value!=int(value) else '#,##0;[Red]-#,##0'
    c.alignment=Alignment(horizontal='right',vertical='center')
   if col in (1,5):c.alignment=Alignment(horizontal='center',vertical='center')
  description=values[1]+' '+values[2]
  ws.row_dimensions[row].height=max(28,min(105,14*(max(len(values[1])//20,len(values[2])//17)+1))) if item else 28
 end=14+count;box(end,1,6,'합계',True,10,'center')
 for col,value in ((7,r['net']),(8,r['vat']),(9,total)):
  c=box(end,col,col,'' if blank else value,True,10,'right');c.number_format='#,##0;[Red]-#,##0'
 ws.row_dimensions[end].height=28
 box(end+2,1,9,'MEMO',True,10);ws.row_dimensions[end+2].height=24
 box(end+3,1,9,r.get('notes',''));ws.row_dimensions[end+3].height=max(62,min(180,16*(r.get('notes','').count('\n')+2)))
 ws.cell(end+3,1).alignment=Alignment(vertical='top',wrap_text=True)
 box(end+5,1,9,'입금계좌  '+' / '.join(supplier[k] for k in ('bank','bankAccount','bankHolder') if supplier.get(k)),False,9,border=False)
 # Strong outer rules complement the visible internal grid, including merged header fields.
 for first,last in ((4,11),(13,end),(end+2,end+3)):
  for row in range(first,last+1):
   for col in range(1,10):
    c=ws.cell(row,col);b=c.border
    c.border=Border(left=strong if col==1 else b.left,right=strong if col==9 else b.right,top=strong if row==first else b.top,bottom=strong if row==last else b.bottom)
 ws.freeze_panes='D14';ws.print_title_rows='1:13';ws.print_area=f'A1:I{end+5}'
 ws.page_setup.orientation='portrait';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
 ws.sheet_properties.pageSetUpPr.fitToPage=True
 ws.page_margins.left=ws.page_margins.right=.25;ws.page_margins.top=ws.page_margins.bottom=.35
 ws.oddFooter.center.text='&P / &N';ws.oddFooter.right.text='Growth ERP'
 out=io.BytesIO();w.save(out);return out.getvalue()


def generate(s,d,workspace):
 kind=d.get('kind');fmt=d.get('format','pdf');ledger.need(kind in ('quote','voucher') and fmt in ('pdf','xlsx'),'문서 종류·형식을 확인해 주세요.')
 record=ledger.get(s,'quotes' if kind=='quote' else 'vouchers',d['id']);r=snapshot(s,kind,d['id'])
 fingerprint=hashlib.sha256(json.dumps({'templateVersion':TEMPLATE_VERSION,**{k:v for k,v in r.items() if k not in ('pdfDocumentId','excelDocumentId')}},sort_keys=True,ensure_ascii=False).encode()).hexdigest()
 cached=next((x for x in s['documents'] if x.get('erpSourceId')==r['id'] and x.get('erpFingerprint')==fingerprint and x.get('format')==fmt),None)
 if cached:
  if kind=='quote':record['excelDocumentId' if fmt=='xlsx' else 'pdfDocumentId']=cached['id']
  return {'recordId':r['id'],'documentId':cached['id']},None
 if fmt=='xlsx':content=render_xlsx(kind,r,workspace=='demo')
 else:
  from pdf_documents import render_pdf
  content=render_pdf('quote' if kind=='quote' else 'delivery',r,workspace=='demo')
 label='견적서' if kind=='quote' else '매입거래명세서' if r['kind']=='purchase' else '거래명세서'
 key=ledger.uid('doc');safe=re.sub(r'[^\w가-힣-]','_',r['number'])[:100]
 doc=dict(id=key,generated=True,name=f'{label}-{safe}'+(f"-v{r['version']}" if kind=='quote' else '')+'.'+fmt,storedName=key+'.'+fmt,sha256=hashlib.sha256(content).hexdigest(),date=ledger.now()[:10],kind=label,status='등록 완료',accountId=r['accountId'],orderId=r.get('orderId',''),quoteId=r['id'] if kind=='quote' else '',voucherId=r['id'] if kind=='voucher' else '',erpSourceId=r['id'],erpFingerprint=fingerprint,format=fmt,createdAt=ledger.now(),snapshot=r)
 s['documents'].append(doc)
 if kind=='quote':record['excelDocumentId' if fmt=='xlsx' else 'pdfDocumentId']=key
 return {'recordId':r['id'],'documentId':key},content
